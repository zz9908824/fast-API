"""Private OCR service for matter-import files. It never exposes files publicly."""

import io
import os
from pathlib import Path
from typing import Any

import fitz
import numpy as np
from docx import Document
from fastapi import FastAPI, HTTPException
from pydantic import BaseModel
from PIL import Image

app = FastAPI(title="Matter import OCR worker")
STORAGE_ROOT = Path(os.getenv("OCR_STORAGE_ROOT", "/data/imports")).resolve()
APPLICATION_STORAGE_ROOT = Path(os.getenv("OCR_APPLICATION_STORAGE_ROOT", "/data/application-documents")).resolve()
DET_MODEL_DIR = os.getenv("OCR_DET_MODEL_DIR", "")
REC_MODEL_DIR = os.getenv("OCR_REC_MODEL_DIR", "")
OCR_CPU_THREADS = max(1, int(os.getenv("OCR_CPU_THREADS", "8")))
_ocr: Any | None = None


class ExtractRequest(BaseModel):
    storage_key: str
    mime_type: str
    storage_area: str = "imports"


def safe_path(storage_key: str, storage_area: str = "imports") -> Path:
    root = APPLICATION_STORAGE_ROOT if storage_area == "applications" else STORAGE_ROOT
    if storage_area not in {"imports", "applications"}:
        raise HTTPException(status_code=400, detail="不支持的文件存储区域")
    target = (root / storage_key).resolve()
    if root not in target.parents:
        raise HTTPException(status_code=400, detail="非法文件路径")
    if not target.is_file():
        raise HTTPException(status_code=404, detail="导入文件不存在")
    return target


def get_ocr() -> Any:
    global _ocr
    if _ocr is None:
        if not Path(DET_MODEL_DIR).is_dir() or not Path(REC_MODEL_DIR).is_dir():
            raise HTTPException(status_code=503, detail="本地 PP-OCRv5 检测或识别模型目录不可用")
        try:
            from paddleocr import PaddleOCR
            _ocr = PaddleOCR(
                text_detection_model_name="PP-OCRv5_server_det",
                text_detection_model_dir=DET_MODEL_DIR,
                text_recognition_model_name="PP-OCRv5_server_rec",
                text_recognition_model_dir=REC_MODEL_DIR,
                # Keep inference fully local: these optional helpers otherwise
                # download extra official models on first use.
                use_doc_orientation_classify=False,
                use_doc_unwarping=False,
                use_textline_orientation=False,
                enable_mkldnn=False,
                cpu_threads=OCR_CPU_THREADS,
            )
        except Exception as error:
            raise HTTPException(status_code=503, detail=f"无法初始化 PaddleOCR：{error}") from error
    return _ocr


def blocks_from_image(image: Image.Image, page_number: int) -> list[dict[str, Any]]:
    """Use the V3 pipeline and retain text plus any returned position/confidence metadata."""
    ocr = get_ocr()
    try:
        results = list(ocr.predict(np.asarray(image.convert("RGB"))))
    except Exception as error:
        raise HTTPException(status_code=422, detail=f"图片 OCR 失败：{error}") from error
    blocks: list[dict[str, Any]] = []
    for result in results:
        payload = result.json if hasattr(result, "json") else result
        if callable(payload):
            payload = payload()
        if isinstance(payload, str):
            import json
            payload = json.loads(payload)
        if isinstance(payload,dict):
            payload = payload.get("res",payload)
        texts = payload.get("rec_texts", []) if isinstance(payload, dict) else []
        scores = payload.get("rec_scores", []) if isinstance(payload, dict) else []
        boxes = payload.get("rec_boxes", []) if isinstance(payload, dict) else []
        for index, value in enumerate(texts):
            if not isinstance(value, str) or not value.strip():
                continue
            blocks.append({"block_id": f"p{page_number}-ocr-{index + 1}", "page_number": page_number,
                           "text": value.strip(), "bbox": boxes[index] if index < len(boxes) else None,
                           "confidence": scores[index] if index < len(scores) else None, "method": "paddleocr"})
    return blocks


def extract_pdf(path: Path) -> list[dict[str, Any]]:
    document = fitz.open(path)
    blocks: list[dict[str, Any]] = []
    for index, page in enumerate(document, start=1):
        text = page.get_text("text").strip()
        if text:
            blocks.append({"block_id": f"p{index}-text", "page_number": index, "text": text,
                           "bbox": None, "confidence": None, "method": "pymupdf"})
            continue
        # Scanned government PDFs can be much slower than born-digital PDFs.
        # 1.25x preserves Chinese character detail while avoiding the 2x raster
        # image (over 2.5 times as many pixels) that caused request timeouts on CPU.
        pixmap = page.get_pixmap(matrix=fitz.Matrix(1.25, 1.25), alpha=False)
        image = Image.open(io.BytesIO(pixmap.tobytes("png"))).convert("RGB")
        blocks.extend(blocks_from_image(image, index))
    return blocks


def extract_docx(path: Path) -> list[dict[str, Any]]:
    document = Document(path)
    pieces = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                pieces.append(row_text)
    return [{"block_id": f"docx-{index}", "page_number": None, "text": value,
             "bbox": None, "confidence": None, "method": "python-docx"} for index, value in enumerate(pieces, start=1)]


@app.on_event("startup")
def warm_ocr_model() -> None:
    """Do model initialization before reporting the service as ready."""
    get_ocr()


@app.get("/health")
def health() -> dict[str, bool]:
    return {"ok": True}


@app.post("/extract")
def extract(request: ExtractRequest) -> dict[str, list[dict[str, Any]]]:
    path = safe_path(request.storage_key, request.storage_area)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        blocks = extract_pdf(path)
    elif suffix == ".docx":
        blocks = extract_docx(path)
    elif suffix == ".txt":
        blocks = [{"block_id": "txt-1", "page_number": None, "text": path.read_text(encoding="utf-8", errors="replace"), "bbox": None, "confidence": None, "method": "text"}]
    elif suffix in {".jpg", ".jpeg", ".png", ".webp"}:
        with Image.open(path) as image:
            blocks = blocks_from_image(image.convert("RGB"), 1)
    else:
        raise HTTPException(status_code=415, detail="不支持的文件类型")
    return {"blocks": [block for block in blocks if block["text"].strip()]}
